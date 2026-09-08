"""Word export for the global HAZOP recommendation catalog."""

from worksheet_word_export import (
    PAPER_SIZES_MM,
    _prevent_row_split,
    _set_cell_borders,
    _set_cell_shading,
    _set_cell_text,
    _set_repeat_table_header,
)


_HEADERS = [
    'Rek. nr', 'Rekommendation', 'Ansvarig', 'Ska vara åtgärdat',
    'Status', 'Referenser (studie.nod.avvikelse.orsak.konsekvens)',
]
_WIDTH_RATIOS = [12, 50, 25, 18, 16, 18]


def _build_position_maps(db):
    node_pos, dev_pos, cause_pos, cons_pos = {}, {}, {}, {}
    dev_node, cause_dev, cons_cause = {}, {}, {}
    for node_number, node in enumerate(db.nodes(), start=1):
        node_pos[node['id']] = node_number
        numbers_by_description = {}
        for deviation in db.deviations(node['id']):
            description = deviation['description'] or ''
            numbers_by_description.setdefault(description, len(numbers_by_description) + 1)
            deviation_number = numbers_by_description[description]
            dev_pos[deviation['id']] = deviation_number
            dev_node[deviation['id']] = node['id']
            for cause_number, cause in enumerate(
                    db.causes_for_deviation(deviation['id']), start=1):
                cause_pos[cause['id']] = cause_number
                cause_dev[cause['id']] = deviation['id']
                for consequence_number, consequence in enumerate(
                        db.consequences(cause['id']), start=1):
                    cons_pos[consequence['id']] = consequence_number
                    cons_cause[consequence['id']] = cause['id']
    return {
        'node_pos': node_pos, 'dev_pos': dev_pos, 'cause_pos': cause_pos,
        'cons_pos': cons_pos, 'dev_node': dev_node, 'cause_dev': cause_dev,
        'cons_cause': cons_cause,
    }


def _reference_for_consequence(consequence_id, maps):
    cause_id = maps['cons_cause'].get(consequence_id)
    consequence_number = maps['cons_pos'].get(consequence_id)
    if cause_id is None or consequence_number is None:
        return None
    deviation_id = maps['cause_dev'].get(cause_id)
    cause_number = maps['cause_pos'].get(cause_id)
    if deviation_id is None or cause_number is None:
        return None
    node_id = maps['dev_node'].get(deviation_id)
    deviation_number = maps['dev_pos'].get(deviation_id)
    node_number = maps['node_pos'].get(node_id)
    if node_id is None or deviation_number is None or node_number is None:
        return None
    return f'1.{node_number}.{deviation_number}.{cause_number}.{consequence_number}'


def _add_recommendation_table(document, rows, page_width_mm, margin_mm):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Mm

    table = document.add_table(rows=1, cols=len(_HEADERS))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable_width = page_width_mm - 2 * margin_mm
    total_ratio = sum(_WIDTH_RATIOS)
    widths = [usable_width * ratio / total_ratio for ratio in _WIDTH_RATIOS]

    for column, header in enumerate(_HEADERS):
        cell = table.cell(0, column)
        cell.width = Mm(widths[column])
        _set_cell_shading(cell, 'EEECE1')
        _set_cell_borders(cell)
        _set_cell_text(cell, header, bold=True, colour='17191C', size=8,
                       align='center')
    _set_repeat_table_header(table.rows[0])
    _prevent_row_split(table.rows[0])

    for values in rows:
        cells = table.add_row().cells
        for column, cell in enumerate(cells):
            cell.width = Mm(widths[column])
            _set_cell_shading(cell, 'FFFFFF')
            _set_cell_borders(cell)
            _set_cell_text(
                cell, values[column], size=8,
                align='center' if column in (0, 3, 4, 5) else 'left')
        _prevent_row_split(table.rows[-1])


def export_recommendations_word(db, filepath, paper_size='A4'):
    """Export the global recommendation catalog to landscape Word."""
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Mm
    except ImportError:
        return False, 'python-docx saknas.\nKör: pip install python-docx'

    paper_size = str(paper_size or 'A4').upper()
    if paper_size not in PAPER_SIZES_MM:
        return False, f'Okänt pappersformat: {paper_size}'
    try:
        maps = _build_position_maps(db)
        rows = [
            [
                f"{int(rec['display_number']):03d}",
                dict(rec).get('description') or '',
                dict(rec).get('responsible') or '',
                dict(rec).get('due_date') or '',
                dict(rec).get('status') or '',
                ', '.join(
                    ref for ref in (
                        _reference_for_consequence(consequence_id, maps)
                        for consequence_id in db.consequences_for_recommendation(
                            rec['id']))
                    if ref is not None) or '—',
            ]
            for rec in db.all_recommendations()
        ]
        page_width, page_height = PAPER_SIZES_MM[paper_size]
        margin = 8 if paper_size == 'A3' else 10

        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(page_width)
        section.page_height = Mm(page_height)
        section.left_margin = Mm(margin)
        section.right_margin = Mm(margin)
        section.top_margin = Mm(margin)
        section.bottom_margin = Mm(margin)
        section.header_distance = Mm(3)
        section.footer_distance = Mm(3)

        _add_recommendation_table(document, rows, page_width, margin)
        document.save(filepath)
        return True, ''
    except Exception as error:
        return False, str(error)
