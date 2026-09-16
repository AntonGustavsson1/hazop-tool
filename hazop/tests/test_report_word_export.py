import os
import sys
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from docx import Document
from docx.oxml.ns import qn
import database
from database import Database, load_matrix
from report_word_export import (
    export_report_word, _matrix_display_values, _report_snapshot,
    collect_report_data, _worksheet_references, _annotated_worksheet_rows,
)
from tor_word_export import collect_tor_data, export_tor_word
from worksheet_export import _worksheet_rows
from worksheet_word_export import export_worksheet_word
from recommendations_word_export import (
    _build_position_maps, _reference_for_consequence, export_recommendations_word,
)


class ReportWordExportTests(unittest.TestCase):
    def setUp(self):
        self.db = Database(':memory:')
        self.folder = tempfile.TemporaryDirectory()
        self.path = Path(self.folder.name) / 'report.docx'
        self.node = self.db.nodes()[0]['id']
        self.deviation = self.db.deviations(self.node)[0]['id']
        self.cause = self.db.add_cause(self.deviation)
        self.db.update_cause(self.cause, description='Ventil felar öppen', likelihood=2)
        self.cons = self.db.add_consequence(self.cause)
        self.db.update_consequence(self.cons, 'Överfyllning', 3)
        self.db.set_config('project_name', 'Projekt ÅÄÖ')
        self.db.set_config('project_client', 'Testkund')
        self.rec = self.db.add_recommendation('Kontrollera ventilen', responsible='Drift',
                                             due_date='2026-10-01', status='Öppen')
        self.db.link_recommendation_to_consequence(self.rec, self.cons)

    def tearDown(self):
        self.db.conn.close()
        self.folder.cleanup()

    def export(self, **kwargs):
        ok, error = export_report_word(self.db, self.path, **kwargs)
        self.assertTrue(ok, error)
        return Document(self.path)

    def export_tor(self):
        path = Path(self.folder.name) / 'tor.docx'
        ok, error = export_tor_word(self.db, path)
        self.assertTrue(ok, error)
        return Document(path), path

    def test_tor_contains_only_planning_inputs_and_keeps_them_adjustable(self):
        participant = self.db.add_participant('Anna', 'Andersson', 'Drift', 'Testkund')
        self.db.add_analysis_session('Planerad dag', '2026-10-01', 'Kontoret')
        self.db.add_standard_deviation('Egen ToR-avvikelse')
        self.db.add_project_custom_field('ToRnummer', '262054-ToR-01')
        self.db.add_project_custom_field('ToRdatum', '2026-09-14')
        self.db.add_project_revision('A', '2026-09-14', 'Första utgåva', 'Anna Andersson')
        data = collect_tor_data(self.db)
        self.assertNotIn('rows', data)
        self.assertNotIn('recommendations', data)
        doc, path = self.export_tor()
        body = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
        table_text = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn('Antal deltagare och närvaro uppdateras vid analysen.', body)
        self.assertIn('Guideorden kan justeras', body)
        self.assertIn('Nodindelningen och markeringarna är planeringsunderlag och kan justeras', body)
        self.assertIn('Riskgraph eller LOPA', body)
        self.assertIn('Anna', table_text)
        self.assertIn('Egen ToR-avvikelse', table_text)
        self.assertNotIn('Kontrollera ventilen', body + table_text)
        self.assertNotIn('Överfyllning', body + table_text)
        self.assertTrue(path.exists())

    def test_tor_includes_planned_risk_scales_and_decision_framework(self):
        """ToR reuses the report's risk definitions without showing results."""
        matrix = deepcopy(self.db.get_risk_matrix() or database.DEFAULT_MATRIX)
        matrix['risk_level_definitions'] = [{
            'color': '#EF4444',
            'label': 'Oacceptabel risk',
            'definition': 'Riskreducering ska övervägas före fortsatt hantering.',
        }]
        self.db.set_risk_matrix(matrix)
        person = self.db.consequence_categories()[0]
        self.db.set_severity_definition(1, person['id'], 'Lindrig personskada')
        self.db.set_severity_definition(2, person['id'], 'Allvarlig personskada')
        self.db.add_project_custom_field(
            'Frekvensunderlag', 'Frekvens bedöms av analysgruppen utifrån driftfall.')
        self.db.add_project_custom_field(
            'Barriärunderlag', 'Skyddets oberoende ska bekräftas av gruppen.')

        doc, _path = self.export_tor()
        headings = [paragraph.text for paragraph in doc.paragraphs
                    if paragraph.style.name in ('Heading 1', 'Heading 2')]
        captions = [paragraph.text for paragraph in doc.paragraphs
                    if paragraph.style.name == 'Caption']
        table_text = '\n'.join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        body = '\n'.join(paragraph.text for paragraph in doc.paragraphs)

        self.assertIn('5.2 Planerade acceptanskriterier', headings)
        self.assertIn('5.3 Frekvensskala', headings)
        self.assertIn('5.4 Konsekvensdefinitioner', headings)
        self.assertIn('5.5 Barriärer, enablers, riskgraph och LOPA', headings)
        for caption in (
            'Tabell 5-2 Planerade acceptanskriterier',
            'Tabell 5-3 Planerade frekvensnivåer och definitioner',
            'Tabell 5-4 Planerade konsekvensdefinitioner',
        ):
            self.assertIn(caption, captions)
        self.assertIn('Oacceptabel risk', table_text)
        self.assertIn('Lindrig personskada', table_text)
        self.assertIn('Frekvens bedöms av analysgruppen', body)
        self.assertIn('Skyddets oberoende ska bekräftas', body)
        self.assertIn('funktion, relevans och oberoende', body)
        self.assertNotIn('Överfyllning', body + table_text)
        self.assertTrue(any(section.page_width > section.page_height
                            for section in doc.sections))

    def test_tor_keeps_frequency_and_consequence_definitions_on_their_storage_axes(self):
        """Changing the visual matrix orientation must not swap the two lists."""
        matrix = deepcopy(self.db.get_risk_matrix() or database.DEFAULT_MATRIX)
        matrix.update({
            'rows': 2,
            'cols': 3,
            # Display consequences horizontally, like the ProSa IPS template.
            'x_axis': 'consequence',
            'x_reversed': False,
            'y_reversed': False,
            # x data always remains the frequency scale in storage.
            'x_codes': ['F0', 'F1', 'F2'],
            'x_labels': ['Frekvens noll', 'Frekvens ett', 'Frekvens två'],
            'y_codes': ['C1', 'C2'],
            'y_labels': ['Konsekvens ett', 'Konsekvens två'],
            'cell_labels': [['Låg', 'Medium', 'Hög'], ['Medium', 'Hög', 'Hög']],
            'cell_colors': [['#22AA66', '#FFFF00', '#EE4444'],
                            ['#FFFF00', '#EE4444', '#EE4444']],
            'cell_fg_colors': [['#000000'] * 3, ['#000000'] * 3],
        })
        self.db.set_risk_matrix(matrix)
        person = self.db.consequence_categories()[0]
        self.db.set_severity_definition(1, person['id'], 'Definition C1')
        self.db.set_severity_definition(2, person['id'], 'Definition C2')

        doc, _path = self.export_tor()
        frequency_table = next(
            table for table in doc.tables
            if table.cell(0, 0).text == 'Nivå' and table.cell(1, 0).text == 'F0')
        consequence_table = next(
            table for table in doc.tables
            if table.cell(0, 0).text == 'Nivå'
            and table.cell(0, 1).text == 'Benämning')
        self.assertEqual(
            [('F0', 'Frekvens noll'), ('F1', 'Frekvens ett'),
             ('F2', 'Frekvens två')],
            [(row.cells[0].text, row.cells[1].text)
             for row in frequency_table.rows[1:]])
        self.assertEqual(
            [('C1', 'Konsekvens ett'), ('C2', 'Konsekvens två')],
            [(row.cells[0].text, row.cells[1].text)
             for row in consequence_table.rows[1:]])

    def test_tor_includes_compact_active_frequency_catalogue_and_formal_basis(self):
        """The ToR must contain the current reusable catalogue, once per cause."""
        data = collect_tor_data(self.db)
        catalogue = data['standard_cause_catalogue']
        self.assertEqual(32, len(catalogue))
        self.assertEqual(
            {'Manuell ventil', 'On-off ventil', 'Reglerventil', 'Instrument',
             'Backventil', 'Säkerhetsventil / sprängbleck', 'Pump',
             'Kompressor / fläkt', 'Filter / sil',
             'Värmeväxlare / kylare / värmare', 'Blandare / omrörare',
             'Operatör / procedur / underhåll'},
            {cause['object_name'] for cause in catalogue})
        self.assertNotIn('Tank / kärl / kolonn',
                         {cause['object_name'] for cause in catalogue})

        doc, _path = self.export_tor()
        catalogue_table = next(
            table for table in doc.tables
            if table.cell(0, 0).text == 'Objekt'
            and table.cell(0, 1).text == 'Föreslagna standardorsaker och frekvenser')
        self.assertEqual(13, len(catalogue_table.rows))
        catalogue_text = '\n'.join(
            cell.text for row in catalogue_table.rows for cell in row.cells)
        body = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
        for value in ('Ventil felaktigt stängd — 0,01/år',
                      'Reglerventil felar öppen — 0,09/år',
                      'Tubläckage — 0,01/år',
                      'Omrörare stopp — 0,1/år'):
            self.assertIn(value, catalogue_text)
        self.assertNotIn('Endoterm reaktion / avdunstning', catalogue_text)
        self.assertIn('mindre än en farlig felfunktion per 100 000 timmar', body)
        self.assertIn('inte användas som en direkt komponentdatauppgift', body)
        self.assertIn('Renare eller smutsigare medier', body)
        self.assertIn('manöverfrekvens', body)
        self.assertIn('elnätets dokumenterade tillförlitlighet', body)
        for standard in ('IEC 61882:2016', 'IEC 61511-1:2016',
                         'IEC 61511-2:2016', 'IEC 61511-3:2017',
                         'IEC 61508-1:2010', 'IEC 61508-2:2010'):
            self.assertIn(standard, body)
        self.assertGreaterEqual(len(doc.sections), 6)
        self.assertLess(doc.sections[-2].page_width, doc.sections[-2].page_height)
        self.assertLess(doc.sections[-1].page_width, doc.sections[-1].page_height)

    def test_tor_catalogue_refreshes_when_a_standard_cause_changes(self):
        """The ToR reads the current catalogue on every export, not a copy."""
        node_type = self.db.node_types()[0]
        manual_valve = next(
            item for item in self.db.standard_objects() if item['name'] == 'Manuell ventil')
        deviation = self.db.standard_deviations_for_node_type(node_type['id'])[0]
        group_id = self.db.add_standard_cause_group(
            node_type['id'], manual_valve['id'], 'Provbar dynamisk orsak', 0.03)
        self.db.set_standard_cause_group_deviation(group_id, deviation['id'], True)

        doc, _path = self.export_tor()
        catalogue_table = next(
            table for table in doc.tables
            if table.cell(0, 0).text == 'Objekt'
            and table.cell(0, 1).text == 'Föreslagna standardorsaker och frekvenser')
        catalogue_text = '\n'.join(
            cell.text for row in catalogue_table.rows for cell in row.cells)
        self.assertIn('Provbar dynamisk orsak — 0,03/år', catalogue_text)

        self.db.update_standard_cause_group(group_id, frequency=0.04)
        doc, _path = self.export_tor()
        catalogue_table = next(
            table for table in doc.tables
            if table.cell(0, 0).text == 'Objekt'
            and table.cell(0, 1).text == 'Föreslagna standardorsaker och frekvenser')
        catalogue_text = '\n'.join(
            cell.text for row in catalogue_table.rows for cell in row.cells)
        self.assertIn('Provbar dynamisk orsak — 0,04/år', catalogue_text)
        self.assertNotIn('Provbar dynamisk orsak — 0,03/år', catalogue_text)

    def test_complete_report_uses_project_and_signoff_fields_without_legacy_customer(self):
        self.db.add_project_custom_field('Framtagen av', 'Författare')
        self.db.add_project_custom_field('Rapportnummer', '262054-Report-02')
        self.db.add_project_custom_field('Omfattning', 'Tank och lastning')
        self.db.add_project_custom_field('Eget fält', 'Eget värde')
        self.export()
        with ZipFile(self.path) as archive:
            xml = '\n'.join(archive.read(n).decode('utf-8') for n in archive.namelist()
                            if n.endswith('.xml'))
        for value in ('Projekt ÅÄÖ', 'Testkund', 'Författare', 'Tank och lastning',
                      '262054-Report-02', 'Bilaga 4', '5.2 Fortsatt hantering'):
            self.assertIn(value, xml)
        for value in ('Mondi', 'Aurora', 'HEAD Engineering', '53 åtgärder', 'SIL 3'):
            self.assertNotIn(value, xml)
        self.assertIn('w:updateFields', xml)
        self.assertIn('TOC ', xml)

    def test_cover_metadata_labels_have_a_space_before_their_values(self):
        self.db.add_project_custom_field('Distribution', 'Enligt kundens anvisning')
        doc = self.export()
        metadata = [
            paragraph.text
            for row in doc.tables[0].rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
        self.assertTrue(any(text.startswith('Titel: HAZOP för ') for text in metadata))
        self.assertTrue(any(text.startswith('Datum: ') for text in metadata))
        self.assertTrue(any(text.startswith('Distribution: Enligt ') for text in metadata))

    def test_protocol_export_repairs_a_missing_space_beside_och_only_in_output(self):
        rows = [{'values': ['barriäroch kontroll', 'ventiler ochstängning']}]
        exported_rows = _annotated_worksheet_rows(self.db, rows)
        self.assertEqual(
            exported_rows[0]['values'],
            ['barriär och kontroll', 'ventiler och stängning'])
        self.assertEqual(rows[0]['values'][0], 'barriäroch kontroll')

    def test_source_backed_prosa_furniture_and_descriptive_chapter_text_are_preserved(self):
        doc = self.export()
        self.assertEqual(doc.styles['Normal'].font.name, 'Aptos')
        self.assertEqual(str(doc.styles['Heading 1'].font.color.rgb), '509628')
        self.assertEqual(str(doc.styles['Heading 2'].font.color.rgb), '509628')
        self.assertTrue(doc.sections[0].different_first_page_header_footer)
        self.assertTrue(all(not section.header.is_linked_to_previous
                            for section in doc.sections[2:]))
        self.assertEqual(
            doc.sections[2]._sectPr.find(qn('w:pgNumType')).get(qn('w:start')),
            '1',
        )
        for section in doc.sections[:2]:
            self.assertNotIn(
                'PAGE',
                [instruction.text or ''
                 for instruction in section.footer._element.iter(qn('w:instrText'))],
            )
        for section in doc.sections[2:]:
            for header in (section.header, section.first_page_header):
                blip = next(header._element.iter(qn('a:blip')), None)
                self.assertIsNotNone(blip)
                image_relation = header.part.rels[blip.get(qn('r:embed'))]
                self.assertEqual(
                    str(image_relation.target_part.partname),
                    '/word/media/image4.wmf',
                    'The generated header must retain the complete ProSa logo '
                    'from the source running header.')
                instructions = [instruction.text or '' for instruction
                                in header._element.iter(qn('w:instrText'))]
                self.assertIn('REF rapportdatum', instructions)
                self.assertIn('REF rapportnummer', instructions)
        bookmark_names = {
            bookmark.get(qn('w:name'))
            for bookmark in doc.element.iter(qn('w:bookmarkStart'))
        }
        self.assertIn('rapportdatum', bookmark_names)
        self.assertIn('rapportnummer', bookmark_names)
        body_text = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
        for text in (
            'P&ID-ritningarna och övriga registrerade dokument för riskanalysen redovisas i tabell 2-1',
            'Datum, tid och plats för genomförda analystillfällen redovisas i tabell 3-1',
            'Tabell B4-1 redovisar varje rekommendation',
            'HAZOP-studien genomfördes som en gemensam, tvärdisciplinär genomgång',
            'Resultatet ska användas som underlag för fortsatt riskhantering av systemet.',
            'Resultatet har dokumenterats i ett gemensamt protokoll som kan användas vid fortsatt riskhantering.',
            'Nodindelningen redovisas i bilaga 2 och tillhörande P&ID-nodmarkeringar redovisas i bilaga 5.',
            'Tabell 4-1 visar den riskmatris som användes i studien. Axelriktning, färger och nivånamn framgår av tabellen.',
            'vid utfärdande av rapport.',
            '3.1 Metod och arbetssätt',
            'Arbetsgången har varit följande',
            'Bilaga 1 Avvikelser',
            '5.1 Resultat',
            '5.2 Fortsatt hantering',
            'Begränsning genom egensäker design',
            'Riskreducering genom procedurer',
            'Bilaga 5 Nodmarkeringar',
        ):
            self.assertIn(text, body_text)
        self.assertNotIn('Bilaga 1 HAZOP metodik', body_text)
        self.assertNotIn('B1.1 Registrerade avvikelser', body_text)
        self.assertNotIn('B4.1 Rekommendationsregister', body_text)
        with ZipFile(self.path) as archive:
            header_xml = archive.read('word/header4.xml').decode('utf-8')
        self.assertIn('HAZOP f', header_xml)

    def test_missing_values_are_yellow_while_zero_is_preserved(self):
        self.db.add_project_custom_field('Rapportrevision', '0')
        doc = self.export()
        highlights = doc.element.xpath('.//w:r[w:rPr/w:highlight[@w:val="yellow"]]')
        text = '\n'.join(''.join(r.itertext()) for r in highlights)
        self.assertIn('[KOMPLETTERA: revisionsdatum]', text)
        self.assertIn('[KOMPLETTERA: granskad av]', text)
        self.assertNotIn('[KOMPLETTERA: rapportrevision]', text)

    def test_snapshot_preserves_database_and_matrix_binding_even_on_failure(self):
        other = Database(':memory:')
        try:
            load_matrix(other)
            matrix_before = database.get_matrix()
            before = self.db.conn.serialize()
            undo_before = list(self.db._undo_stack)
            with patch('report_word_export.build_report', side_effect=ValueError('forced')):
                ok, error = export_report_word(self.db, self.path)
            self.assertFalse(ok)
            self.assertEqual(error, 'forced')
            self.assertEqual(self.db.conn.serialize(), before)
            self.assertEqual(self.db._undo_stack, undo_before)
            self.assertIs(database._risk_matrix_cache._db, other)
            self.assertIs(database.get_matrix(), matrix_before)
        finally:
            other.conn.close()

    def test_export_from_wal_source_is_read_only_and_not_a_migration(self):
        path = Path(self.folder.name) / 'project.db'
        disk = Database(path)
        try:
            before = disk.conn.serialize()
            ok, error = export_report_word(disk, self.path)
            self.assertTrue(ok, error)
            self.assertEqual(before, disk.conn.serialize())
        finally:
            disk.conn.close()

    def test_attendance_distinguishes_missing_absent_and_present_with_notes(self):
        first = self.db.add_participant('Anna', 'Andersson')
        second = self.db.add_participant('Bo', 'Berg')
        third = self.db.add_participant('Carl', 'Carlsson')
        session = self.db.add_analysis_session('Dag 1', '2026-09-08', 'Testplats')
        self.db.set_attendance(first, session, False)
        self.db.set_attendance(second, session, True)
        self.db.set_attendance_note(second, session, 'Del av dagen')
        doc = self.export()
        table = next(t for t in doc.tables if t.cell(0, 0).text == 'Deltagare')
        self.assertEqual(table.cell(1, 1).text, 'Frånvarande')
        self.assertEqual(table.cell(2, 1).text, 'Närvarande\nDel av dagen')
        self.assertEqual(table.cell(3, 1).text, '[KOMPLETTERA: närvaro]')

    def test_participant_company_is_a_standard_report_column(self):
        self.db.add_participant('Anna', 'Andersson', company='ProSa')
        doc = self.export()
        table = next(t for t in doc.tables if t.cell(0, 0).text == 'Förnamn')
        self.assertEqual(
            [cell.text for cell in table.rows[0].cells],
            ['Förnamn', 'Efternamn', 'Företag', 'Roll och övriga deltagaruppgifter'],
        )
        self.assertEqual(table.cell(1, 2).text, 'ProSa')
        body_text = '\n'.join(paragraph.text for paragraph in doc.paragraphs)
        self.assertNotIn('Anteckningar kan användas för', body_text)

    def test_duplicate_guide_words_follow_worksheet_numbering_in_both_reports(self):
        equipment = self.db.add_equipment_item('V-200', 'V-200', 'V', 0, 'Ventil', '', False)
        deviation = self.db.conn.execute(
            'INSERT INTO deviations(node_id,description,equipment_id) VALUES (?,?,?)',
            (self.node, self.db.deviations(self.node)[0]['description'], equipment)).lastrowid
        self.db.conn.commit()
        cause = self.db.add_cause(deviation)
        self.db.update_cause(cause, 'Andra orsaken', 2)
        consequence = self.db.add_consequence(cause)
        self.db.update_consequence(consequence, 'Andra konsekvensen', 2)
        self.db.link_recommendation_to_consequence(self.rec, consequence)
        with _report_snapshot(self.db) as snap:
            data = collect_report_data(snap)
            reference = data['refs'][consequence]
            maps = _build_position_maps(snap)
            self.assertEqual(reference, _reference_for_consequence(consequence, maps))
            self.assertIn(reference, data['recommendation_rows'][0][5])

    def test_appendix_reuses_all_worksheet_values_and_recommendation_numbers(self):
        with _report_snapshot(self.db) as snap:
            rows = list(_worksheet_rows(snap))
        doc = self.export()
        worksheet_tables = [t for t in doc.tables if t.cell(0, 0).text == 'Nod']
        all_values = '\n'.join(c.text for t in worksheet_tables for r in t.rows for c in r.cells)
        for row in rows:
            for value in row['values']:
                if value:
                    self.assertIn(value, all_values)
        rec_table = next(t for t in doc.tables if t.cell(0, 0).text == 'Rek. nr')
        self.assertEqual(len(rec_table.rows), 2)
        self.assertEqual(rec_table.cell(1, 0).text, '001')
        self.assertIn('w:tblHeader', rec_table.rows[0]._tr.xml)
        self.assertAlmostEqual(doc.sections[-1].page_width.mm, 420, places=1)

    def test_a4_landscape_and_unknown_format_does_not_overwrite_existing_file(self):
        doc = self.export(paper_size='A4')
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=0.2)
        self.assertAlmostEqual(doc.sections[-1].page_width.mm, 297, places=1)
        original = self.path.read_bytes()
        ok, _ = export_report_word(self.db, self.path, 'A2')
        self.assertFalse(ok)
        self.assertEqual(self.path.read_bytes(), original)

    def test_empty_template_has_no_demo_risk_matrix_or_sil_ratings(self):
        self.db.conn.execute('DELETE FROM nodes')
        self.db.conn.commit()
        doc = self.export(standard_template=True)
        body_text = '\n'.join(p.text for p in doc.paragraphs)
        table_text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertFalse(any(t.cell(0, 0).text == 'Konsekvens / frekvens' for t in doc.tables))
        self.assertTrue(any('studiens riskmatris' in p.text for p in doc.paragraphs))
        self.assertNotIn('noder och designavsikt', body_text)
        self.assertIn('Datum / analystillfälle', table_text)
        self.assertIn('[KOMPLETTERA: antal analystillfällen, noder och rekommendationer]', body_text)
        self.assertNotIn('Sammanställning av huvudresultat', body_text)

    def test_summary_reports_sessions_and_lists_nodes_in_numbered_style(self):
        self.db.add_analysis_session('Tillfälle 1', '2026-09-09', 'Kontoret')
        doc = self.export()
        body_text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Studien genomfördes vid 1 analystillfälle', body_text)
        self.assertIn('Studerade noder', body_text)
        self.assertIn('Analysen omfattade 1 nod.', body_text)
        self.assertNotIn('Analysgruppen dokumenterade', body_text)
        node_names = {row['name'] for row in self.db.nodes()}
        numbered = [p for p in doc.paragraphs if p.style.name == 'List Number']
        self.assertTrue(node_names.issubset({p.text for p in numbered}))
        node_sequence = next(p for p in numbered if p.text in node_names)
        method_sequence = next(p for p in numbered if p.text.startswith('Nodens avsedda'))
        self.assertNotEqual(node_sequence._p.pPr.numPr.numId.val,
                            method_sequence._p.pPr.numPr.numId.val)

    def test_report_structure_is_compact_and_appendix_nodes_are_explicit(self):
        doc = self.export()
        body_text = '\n'.join(p.text for p in doc.paragraphs)
        headings = [p.text for p in doc.paragraphs
                    if p.style.name in ('Heading 1', 'Heading 2')]
        self.assertIn('3 Genomförande', headings)
        self.assertIn('3.3 Analysgrupp och närvaro', headings)
        self.assertNotIn('3.4 Närvaro', headings)
        self.assertNotIn('2.1 Ritningsunderlag', headings)
        self.assertNotIn('B1.1 Registrerade avvikelser', headings)
        self.assertNotIn('B4.1 Rekommendationsregister', headings)
        self.assertIn('B3.1 Ny nod', body_text)
        node_heading = next(p for p in doc.paragraphs if p.text == 'B3.1 Ny nod')
        self.assertEqual(node_heading.style.name, 'Appendix Node Heading')

    def test_protocol_starts_each_node_after_the_first_on_a_new_page(self):
        second_node = self.db.add_node()
        self.db.update_node(second_node, 'Nod två', '', '')
        second_deviation = self.db.deviations(second_node)[0]['id']
        second_cause = self.db.add_cause(second_deviation)
        self.db.update_cause(second_cause, description='Pump stannar', likelihood=2)
        second_consequence = self.db.add_consequence(second_cause)
        self.db.update_consequence(second_consequence, 'Processavbrott', 3)
        self.db.commit()

        doc = self.export()
        paragraphs = list(doc.paragraphs)
        second_heading_index = next(
            index for index, paragraph in enumerate(paragraphs)
            if paragraph.text == 'B3.2 Nod två')

        self.assertIn('w:type="page"', paragraphs[second_heading_index - 1]._p.xml)

    def test_table_captions_and_prose_references_use_word_fields(self):
        self.db.add_reduction_factor(self.cons, 'Nivåmätning', 10)
        doc = self.export()
        captions = [p for p in doc.paragraphs
                    if p.style.name == 'Caption' and p.text.startswith('Tabell ')]
        caption_texts = [p.text for p in captions]
        self.assertIn('Tabell 4-1 Riskmatris', caption_texts)
        self.assertIn('Tabell 4-2 Acceptanskriterier', caption_texts)
        self.assertIn('Tabell 4-3 Frekvensnivåer och definitioner', caption_texts)
        self.assertIn('Tabell 4-5 Typer av använda enablers och RRF', caption_texts)
        self.assertNotIn('Tabell 4-1a Acceptanskriterier', caption_texts)
        standard_template = self.export(standard_template=True)
        template_captions = [
            paragraph.text for paragraph in standard_template.paragraphs
            if paragraph.style.name == 'Caption' and paragraph.text.startswith('Tabell ')
        ]
        for caption in (
            'Tabell 4-1 Riskmatris',
            'Tabell 4-2 Acceptanskriterier',
            'Tabell 4-3 Frekvensnivåer och definitioner',
            'Tabell 4-4 Konsekvensdefinitioner',
        ):
            self.assertIn(caption, template_captions)
        for paragraph in captions:
            instructions = [
                field.text or ''
                for field in paragraph._p.iter(qn('w:instrText'))
            ]
            self.assertTrue(any(instruction.startswith('SEQ Tabell ') for instruction in instructions),
                            paragraph.text)
            self.assertTrue(any(True for _ in paragraph._p.iter(qn('w:bookmarkStart'))),
                            paragraph.text)

        bookmark_names = {
            bookmark.get(qn('w:name'))
            for bookmark in doc.element.iter(qn('w:bookmarkStart'))
        }
        ref_instructions = [
            field.text or ''
            for field in doc.element.iter(qn('w:instrText'))
            if (field.text or '').startswith('REF tabell_')
        ]
        self.assertTrue(ref_instructions)
        self.assertIn('REF tabell_4_1 \\h', ref_instructions)
        self.assertNotIn('REF tabell_5_1 \\h', ref_instructions)
        for instruction in ref_instructions:
            self.assertIn(instruction.split()[1], bookmark_names)
        self.assertFalse(doc.styles['Appendix Node Heading'].paragraph_format.page_break_before)
        headings = [p.text for p in doc.paragraphs
                    if p.style.name in ('Heading 1', 'Heading 2')]
        self.assertIn('Bilaga 5 Nodmarkeringar', headings)
        with ZipFile(self.path) as archive:
            document_xml = archive.read('word/document.xml').decode('utf-8')
        self.assertIn('TOC \\o "1-2"', document_xml)

    def test_document_subheadings_are_added_when_other_documents_exist(self):
        self.db.add_project_custom_field('Övriga referensdokument', 'Driftinstruktion 100')
        doc = self.export()
        headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 2']
        self.assertIn('2.1 Ritningsunderlag', headings)
        self.assertIn('2.2 Övrigt dokumentunderlag', headings)

    def test_report_matrix_uses_same_axis_directions_as_program(self):
        matrix = {
            'rows': 2, 'cols': 3,
            'x_codes': ['F0', 'F1', 'F2'],
            'y_codes': ['C1', 'C2'],
            'cell_labels': [['00', '01', '02'], ['10', '11', '12']],
            'x_axis': 'consequence', 'x_reversed': False, 'y_reversed': False,
        }
        headers, rows, horizontal, vertical, x_frequency = _matrix_display_values(matrix)
        self.assertFalse(x_frequency)
        self.assertEqual(headers, ['Frekvens / konsekvens', 'C1', 'C2'])
        self.assertEqual(vertical, [2, 1, 0])
        self.assertEqual(rows, [
            ['F2', '02', '12'], ['F1', '01', '11'], ['F0', '00', '10'],
        ])

        matrix.update(x_axis='frequency', x_reversed=True, y_reversed=True)
        headers, rows, horizontal, vertical, x_frequency = _matrix_display_values(matrix)
        self.assertTrue(x_frequency)
        self.assertEqual(headers, ['Konsekvens / frekvens', 'F2', 'F1', 'F0'])
        self.assertEqual(vertical, [0, 1])
        self.assertEqual(rows, [['C1', '02', '01', '00'], ['C2', '12', '11', '10']])


class ReportMenuTests(unittest.TestCase):
    def test_tor_menu_cancel_success_and_error(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        import hazop
        fake = SimpleNamespace(db=object(), status_bar=Mock())
        with patch('hazop.QFileDialog.getSaveFileName', return_value=('', '')), \
                patch('tor_word_export.export_tor_word') as export:
            hazop.MainWindow._export_tor(fake)
            export.assert_not_called()
        with patch('hazop.QFileDialog.getSaveFileName', return_value=('tor', '')), \
                patch('hazop.QApplication.focusWidget', return_value=None), \
                patch('hazop.QMessageBox.information'), \
                patch('hazop.QMessageBox.critical') as critical, \
                patch('hazop.QDesktopServices.openUrl') as open_url, \
                patch('tor_word_export.export_tor_word', return_value=(True, '')) as export:
            hazop.MainWindow._export_tor(fake)
            export.assert_called_once_with(fake.db, 'tor.docx')
            open_url.assert_called_once()
            export.return_value = (False, 'locked')
            open_url.reset_mock()
            hazop.MainWindow._export_tor(fake)
            critical.assert_called_once_with(fake, 'Fel vid ToR-export', 'locked')
            open_url.assert_not_called()

    def test_report_menu_cancel_success_and_error(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        import hazop
        fake = SimpleNamespace(db=object(), status_bar=Mock())
        with patch('hazop.QFileDialog.getSaveFileName', return_value=('', '')), \
                patch('report_word_export.export_report_word') as export:
            hazop.MainWindow._export_word_report(fake)
            export.assert_not_called()
        with patch('hazop.QFileDialog.getSaveFileName', return_value=('report', '')), \
                patch('hazop.QInputDialog.getItem', return_value=('A3', True)), \
                patch('hazop.QApplication.focusWidget', return_value=None), \
                patch('hazop.QMessageBox.information'), \
                patch('hazop.QMessageBox.critical') as critical, \
                patch('hazop.QDesktopServices.openUrl') as open_url, \
                patch('report_word_export.export_report_word', return_value=(True, '')) as export:
            hazop.MainWindow._export_word_report(fake)
            export.assert_called_once_with(fake.db, 'report.docx', 'A3')
            open_url.assert_called_once()
            export.return_value = (False, 'locked')
            open_url.reset_mock()
            hazop.MainWindow._export_word_report(fake)
            critical.assert_called_once_with(fake, 'Fel vid rapportexport', 'locked')
            open_url.assert_not_called()


if __name__ == '__main__':
    unittest.main()
