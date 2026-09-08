import os
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

_TEST_DIR = Path(__file__).resolve().parent
_HAZOP_DIR = _TEST_DIR.parent
for _path in (_HAZOP_DIR, _TEST_DIR):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

from database import Database
from worksheet_word_export import export_worksheet_word


class WorksheetWordExportTests(unittest.TestCase):
    def setUp(self):
        self.db = Database(':memory:')
        self.system_id = self.db.conn.execute(
            'SELECT id FROM systems ORDER BY id LIMIT 1').fetchone()[0]
        self.db.conn.execute('DELETE FROM nodes')
        self.db.conn.commit()

    def tearDown(self):
        self.db.conn.close()

    def _make_node(self, name, sort_order):
        return self.db.conn.execute(
            'INSERT INTO nodes(name, system_id, sort_order) VALUES (?,?,?)',
            (name, self.system_id, sort_order)).lastrowid

    def _make_populated_node(self, name, sort_order):
        node_id = self._make_node(name, sort_order)
        deviation_id = self.db.conn.execute(
            'INSERT INTO deviations(node_id, description, sort_order) '
            'VALUES (?,?,?)', (node_id, 'Högt flöde', 1)).lastrowid
        cause_id = self.db.add_cause(deviation_id)
        self.db.update_cause(cause_id, description='Ventil felar', likelihood=2)
        consequence_id = self.db.add_consequence(cause_id)
        self.db.update_consequence(
            consequence_id, description='Övertryck', severity=3,
            category='Process')
        first = self.db.add_safeguard(consequence_id)
        second = self.db.add_safeguard(consequence_id)
        self.db.update_safeguard(first, description='Stäng ventil', rrf=10)
        self.db.update_safeguard(second, description='Larma operatör', rrf=1)
        self.db.conn.commit()

    def test_word_export_uses_worksheet_merges_and_node_page_breaks(self):
        self._make_populated_node('Nod A', 1)
        self._make_populated_node('Nod B', 2)
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            ok, error = export_worksheet_word(self.db, path, paper_size='A4')
            self.assertTrue(ok, error)

            from docx import Document
            document = Document(path)
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(document.tables[0].cell(0, 0).text, 'Nod')
            self.assertEqual(
                document.tables[0].cell(0, 10).text, 'Rekommendation')
            self.assertTrue(
                '<w:tblHeader' in document.tables[0].rows[0]._tr.xml)
            self.assertIn('<w:vMerge', document.tables[0]._tbl.xml)
            merged_cause = document.tables[0].cell(1, 2).text.strip()
            self.assertEqual(merged_cause, '1. Ventil felar')
            self.assertEqual(merged_cause.count('Ventil felar'), 1)

            with ZipFile(path) as archive:
                xml = archive.read('word/document.xml').decode('utf-8')
            self.assertEqual(xml.count('w:type="page"'), 1)
            self.assertNotIn('HAZOP Scenario', xml)
        finally:
            os.unlink(path)

    def test_word_export_supports_landscape_a3(self):
        self._make_populated_node('Nod A', 1)
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            ok, error = export_worksheet_word(self.db, path, paper_size='A3')
            self.assertTrue(ok, error)
            from docx import Document
            section = Document(path).sections[0]
            self.assertGreater(section.page_width, section.page_height)
            self.assertGreater(section.page_width, 5000000)
        finally:
            os.unlink(path)


if __name__ == '__main__':
    unittest.main()
