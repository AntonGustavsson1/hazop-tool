import os
import sys
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
_TEST_DIR = Path(__file__).resolve().parent
_HAZOP_DIR = _TEST_DIR.parent
for _path in (_HAZOP_DIR, _TEST_DIR):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

from database import Database  # noqa: E402
from recommendations_word_export import export_recommendations_word  # noqa: E402


class RecommendationsWordExportTests(unittest.TestCase):
    def setUp(self):
        self.db = Database(':memory:')

    def tearDown(self):
        self.db.conn.close()

    def test_export_contains_catalog_and_repeating_header(self):
        first = self.db.add_recommendation(
            description='Kontrollera ventilen', responsible='Drift',
            due_date='2026-10-01', status='Öppen')
        self.db.add_recommendation(description='Dokumentera kontrollen')
        node_id = self.db.add_node()
        deviation_id = self.db.deviations(node_id)[0]['id']
        cause_id = self.db.add_cause(deviation_id)
        consequence_id = self.db.add_consequence(cause_id)
        self.db.link_recommendation_to_consequence(first, consequence_id)
        node_number = next(
            number for number, node in enumerate(self.db.nodes(), start=1)
            if node['id'] == node_id)
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            ok, error = export_recommendations_word(
                self.db, path, paper_size='A4')
            self.assertTrue(ok, error)

            from docx import Document
            document = Document(path)
            self.assertEqual(len(document.tables), 1)
            table = document.tables[0]
            self.assertEqual(table.cell(0, 0).text, 'Rek. nr')
            self.assertIn('Kontrollera ventilen', table.cell(1, 1).text)
            self.assertIn(f'1.{node_number}.1.1.1', table.cell(1, 5).text)
            self.assertIn('w:fill="EEECE1"', table.cell(0, 0)._tc.xml)
            self.assertIn('<w:tblHeader', table.rows[0]._tr.xml)
        finally:
            os.unlink(path)

    def test_export_supports_landscape_a3(self):
        self.db.add_recommendation(description='Test')
        fd, path = tempfile.mkstemp(suffix='.docx')
        os.close(fd)
        try:
            ok, error = export_recommendations_word(
                self.db, path, paper_size='A3')
            self.assertTrue(ok, error)
            from docx import Document
            section = Document(path).sections[0]
            self.assertGreater(section.page_width, section.page_height)
            self.assertGreater(section.page_width, 5000000)
        finally:
            os.unlink(path)


if __name__ == '__main__':
    unittest.main()
