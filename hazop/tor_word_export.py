"""Terms of Reference export for a planned HAZOP study.

Unlike the report export this module deliberately reads preparation data only.
It never builds worksheet rows or recommendation/result summaries.
"""

from contextlib import contextmanager
from copy import deepcopy
from pathlib import Path
import os
import sqlite3
import tempfile

import database
from report_branding import apply_report_fonts, new_report_document
from report_word_export import (
    MISSING_PREFIX, _bookmark_cover_value, _caption, _field_run,
    _highlight_document, _link_header_value_to_cover, _matrix_display_values,
    _page_setup, _report_snapshot, _table, _value, missing,
)


TOR_FIELDS = (
    'ToRnummer', 'ToRdatum', 'ToRrevision', 'Distribution', 'Bakgrund',
    'Syfte', 'Omfattning', 'Avgränsningar', 'Driftfall',
    'Analysförutsättningar', 'Övriga referensdokument', 'Metodreferens',
    'Riskacceptanskriterier', 'Frekvensunderlag', 'Barriärunderlag',
    'Kontaktperson kund', 'Kontaktuppgifter kund', 'Kontorsadress ProSa',
    'Kontaktuppgifter ProSa', 'Kundadress', 'Uppdragsansvarig',
)
_TOR_DATE_BOOKMARK = 'tordatum'
_TOR_NUMBER_BOOKMARK = 'tornummer'


def collect_tor_data(db):
    """Collect planning inputs without touching analysed scenarios or results."""
    grouped_fields = {}
    for item in db.project_custom_fields():
        key = (item['name'] or '').strip().casefold()
        if key:
            grouped_fields.setdefault(key, []).append(str(item['value'] or '').strip())

    def field(label):
        values = list(dict.fromkeys(
            value for value in grouped_fields.get(label.casefold(), []) if value))
        if len(values) > 1:
            return missing(f'{label} har flera värden: ' + ' / '.join(values))
        return values[0] if values else ''

    node_types = [dict(row) for row in db.node_types()]
    default_node_type = node_types[0] if node_types else None
    standard_cause_catalogue = (
        db.standard_cause_catalogue_with_frequency(default_node_type['id'])
        if default_node_type else [])
    return {
        'field': field,
        'nodes': [dict(row) for row in db.nodes()],
        'sheets': [dict(row) for row in db.get_sheets()],
        'participants': [dict(row) for row in db.list_participants()],
        'participant_columns': db.list_participant_columns(),
        'participant_values': db.get_participant_column_values(),
        'sessions': [dict(row) for row in db.list_analysis_sessions()],
        'guide_words': [dict(row) for row in db.standard_deviations()],
        'standard_cause_catalogue': standard_cause_catalogue,
        'revisions': db.project_revisions(),
        'matrix': database._normalise_matrix(deepcopy(database.get_matrix())),
    }


def _front_heading(document, text):
    from docx.shared import Pt, RGBColor
    paragraph = document.add_paragraph()
    paragraph.style = document.styles['Normal']
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(91, 145, 35)
    return paragraph


def _chapter(document, title):
    from docx.enum.section import WD_SECTION_START
    document.add_section(WD_SECTION_START.NEW_PAGE)
    return document.add_heading(title, 1)


def _copy_running_headers(document, source_elements, source_relationships):
    """Retain the complete source header, including its image relationships."""
    from docx.oxml.ns import qn

    def copy_header(header):
        relationship_ids = {}
        for old_id, relation in source_relationships:
            if relation.is_external:
                new_id = header.part.relate_to(
                    relation.target_ref, relation.reltype, is_external=True)
            else:
                new_id = header.part.relate_to(relation.target_part, relation.reltype)
            relationship_ids[old_id] = new_id
        for child in list(header._element):
            header._element.remove(child)
        for child in deepcopy(source_elements):
            header._element.append(child)
        for element in header._element.iter():
            for attribute in (qn('r:embed'), qn('r:id'), qn('r:link')):
                previous = element.get(attribute)
                if previous in relationship_ids:
                    element.set(attribute, relationship_ids[previous])

    for section in document.sections[2:]:
        section.different_first_page_header_footer = True
        for header in (section.first_page_header, section.header):
            header.is_linked_to_previous = True
            header.is_linked_to_previous = False
            copy_header(header)


def _add_participant_table(document, data):
    document.add_heading('3.1 Preliminära HAZOP-deltagare', 2)
    document.add_paragraph(
        'Tabell 3-1 redovisar de deltagare som är planerade för workshopen. '
        'Rätt kompetenser ska finnas representerade inom process, drift, underhåll '
        'och instrumentering. Antal deltagare och närvaro uppdateras vid analysen.')
    rows = []
    for participant in data['participants']:
        details = []
        if participant.get('role'):
            details.append('Roll: ' + participant['role'])
        for column in data['participant_columns']:
            value = data['participant_values'].get((participant['id'], column['id']))
            if value:
                details.append(f"{column['name']}: {value}")
        rows.append([
            _value(participant.get('first_name'), 'förnamn'),
            _value(participant.get('last_name'), 'efternamn'),
            _value(participant.get('company'), 'företag'),
            '\n'.join(details) or missing('roll eller disciplin'),
        ])
    _caption(document, '3-1', 'Preliminära HAZOP-deltagare')
    _table(document,
           ['Förnamn', 'Efternamn', 'Företag', 'Roll och övriga deltagaruppgifter'],
           rows or [[missing('preliminära deltagare'), '', '', '']],
           [28, 32, 42, 58])


def _add_planned_sessions(document, data):
    document.add_heading('3.2 Planerade analystillfällen', 2)
    document.add_paragraph(
        'Planerade analystillfällen redovisas i tabell 3-2. Tidplanen kan '
        'justeras när underlag, nodindelning och deltagartillgänglighet har bekräftats.')
    rows = []
    for index, session in enumerate(data['sessions'], 1):
        time = '–'.join(value for value in (
            (session.get('start_time') or '').strip(),
            (session.get('end_time') or '').strip()) if value)
        rows.append([str(index), _value(session.get('date'), 'datum'),
                     time or missing('tid'), _value(session.get('location'), 'plats')])
    _caption(document, '3-2', 'Planerade analystillfällen')
    _table(document, ['Tillfälle', 'Datum', 'Tid', 'Plats'],
           rows or [['1', missing('planerat analystillfälle'), '', '']],
           [17, 28, 43, 72])


def _add_nodes(document, db, data):
    document.add_heading('4.2 Preliminär nodindelning och P&ID-markeringar', 2)
    document.add_paragraph(
        'Tabell 4-2 redovisar de noder som för närvarande är definierade i '
        'projektet och om en P&ID-sida har kopplats till noden. Nodindelningen '
        'och markeringarna är planeringsunderlag och kan justeras under workshopen.')
    sheets = {sheet['physical_page']: sheet for sheet in data['sheets']}
    rows = []
    for index, node in enumerate(data['nodes'], 1):
        pages = db.analysis_pages_for_node(node['id'])
        references = []
        for page in pages:
            sheet = sheets.get(page, {})
            drawing = (sheet.get('drawing_number') or sheet.get('sheet_name') or
                       f'PDF-sida {page + 1}')
            references.append(f'{drawing} (sida {page + 1})')
        marker = 'Ja' if pages else 'Nej'
        rows.append([
            str(index), _value(node.get('name'), 'nodnamn'),
            node.get('design_intent') or missing('designavsikt'),
            '\n'.join(references) or node.get('pid_ref') or missing('P&ID-referens'),
            marker,
        ])
    _caption(document, '4-2', 'Preliminär nodindelning och P&ID-markering')
    _table(document, ['Nr', 'Nod', 'Designavsikt', 'P&ID-referens', 'Markerad'],
           rows or [[missing('noder'), '', '', '', '']], [12, 40, 62, 45, 18])


def _add_risk_framework(document, db, data):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    _chapter(document, '5 Riskbedömning och riskreduktion')
    document.add_paragraph(
        'Riskmatrisen och tillhörande definitioner utgör analysgruppens planerade '
        'bedömningsram. De anger inte några resultat. Riskgraph eller LOPA kan '
        'användas vid behov av fördjupad bedömning av identifierade SIF:ar.')
    document.add_heading('5.1 Planerad riskmatris', 2)
    headers, rows, horizontal, vertical, x_frequency = _matrix_display_values(data['matrix'])
    _caption(document, '5-1', 'Planerad riskmatris')
    table = document.add_table(rows=2, cols=len(headers))
    table.autofit = False
    widths = [32] + [18] * (len(headers) - 1)
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 0).text = 'Konsekvens' if x_frequency else 'Frekvens'
    table.cell(0, 1).merge(table.cell(0, len(headers) - 1))
    table.cell(0, 1).text = 'Frekvens' if x_frequency else 'Konsekvens'
    for col, value in enumerate(headers[1:], 1):
        table.cell(1, col).text = str(value)
    for row_number, values in enumerate(rows, 2):
        cells = table.add_row().cells
        cells[0].text = str(values[0])
        for col, value in enumerate(values[1:], 1):
            cells[col].text = str(value)
            consequence, frequency = ((vertical[row_number - 2], horizontal[col - 1])
                                      if x_frequency else (horizontal[col - 1], vertical[row_number - 2]))
            color = data['matrix'].get('cell_colors', [])[consequence][frequency]
            if color:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), str(color).lstrip('#'))
                cells[col]._tc.get_or_add_tcPr().append(shading)
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.alignment = 1
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    for row in table.rows[:2]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.alignment = 1
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(8)

    levels = data['matrix'].get('risk_level_definitions') or []
    if not levels:
        # Older matrix configurations can have coloured cells without an
        # explicit legend. Keep the ToR useful by deriving each distinct,
        # named level without turning a planned assessment into a result.
        by_color = {}
        for row_colors, row_labels in zip(data['matrix'].get('cell_colors', []),
                                          data['matrix'].get('cell_labels', [])):
            for color, label in zip(row_colors, row_labels):
                color, label = str(color or '').strip(), str(label or '').strip()
                if color and label and color not in by_color:
                    by_color[color] = label
        levels = [
            {'color': color, 'label': label, 'definition': ''}
            for color, label in by_color.items()
        ]
    document.add_heading('5.2 Planerade acceptanskriterier', 2)
    document.add_paragraph(
        'Tabell 5-2 redovisar de acceptanskriterier som är planerade för '
        'studien. De ska användas som stöd för analysgruppens bedömning och '
        'för att avgöra när fortsatt utredning eller riskreducering behöver '
        'övervägas. De redovisar inte några bedömda risknivåer.')
    _caption(document, '5-2', 'Planerade acceptanskriterier')
    if levels:
        level_table = _table(document, ['Risknivå', 'Definition'], [
            [_value(item.get('label'), 'risknivå'),
             _value(item.get('definition'), 'acceptanskriterium')]
            for item in levels
        ], [56, 104])
        for row, item in zip(level_table.rows[1:], levels):
            color = str(item.get('color') or '').lstrip('#')
            if color:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), color)
                row.cells[0]._tc.get_or_add_tcPr().append(shading)
                try:
                    red, green, blue = (int(color[index:index + 2], 16)
                                        for index in (0, 2, 4))
                    foreground = ('000000' if (0.299 * red + 0.587 * green +
                                               0.114 * blue) > 160 else 'FFFFFF')
                    for paragraph in row.cells[0].paragraphs:
                        for run in paragraph.runs:
                            run.font.color.rgb = RGBColor.from_string(foreground)
                except (TypeError, ValueError):
                    pass
    else:
        document.add_paragraph(missing('planerade acceptanskriterier'))
    if data['field']('Riskacceptanskriterier'):
        document.add_paragraph(data['field']('Riskacceptanskriterier'))

    # Storage axes are fixed: x is frequency and y is consequence.  ``x_axis``
    # changes only how the risk matrix is drawn, not the meaning of the saved
    # code/label arrays.  Swapping these arrays when the visual axes were
    # reversed made the separate frequency and consequence tables incorrect.
    frequency_codes = data['matrix']['x_codes']
    frequency_labels = data['matrix']['x_labels']
    consequence_codes = data['matrix']['y_codes']
    consequence_labels = data['matrix']['y_labels']
    document.add_heading('5.3 Frekvensskala', 2)
    document.add_paragraph(
        'Tabell 5-3 redovisar de frekvensnivåer och definitioner som är '
        'planerade för HAZOP-bedömningen. Definitionerna ger gruppen en '
        'gemensam utgångspunkt när sannolikheten för ett scenario diskuteras.')
    _caption(document, '5-3', 'Planerade frekvensnivåer och definitioner')
    _table(document, ['Nivå', 'Definition'], [
        [str(code), _value(label, 'frekvensdefinition')]
        for code, label in zip(frequency_codes, frequency_labels)
    ] or [[missing('frekvensnivåer'), '']], [25, 135])
    if data['field']('Frekvensunderlag'):
        document.add_paragraph(data['field']('Frekvensunderlag'))

    from docx.enum.section import WD_SECTION_START
    landscape_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _page_setup(landscape_section, landscape=True, paper='A4')
    document.add_heading('5.4 Konsekvensdefinitioner', 2)
    document.add_paragraph(
        'Konsekvenser bedöms för de kategorier som ingår i studiens '
        'riskmatris. Tabell 5-4 redovisar de definitioner som analysgruppen '
        'ska använda för att skapa en jämförbar bedömning mellan scenarier.')
    categories = [dict(category) for category in db.consequence_categories()]
    definitions = db.get_severity_definitions()
    active_categories = []
    for category in categories:
        values = [definitions.get(index + 1, {}).get(category['id'], '')
                  for index in range(len(consequence_codes))]
        if any((value or '').strip() for value in values):
            active_categories.append(category)
    consequence_rows = []
    for index, code in enumerate(consequence_codes):
        values = [definitions.get(index + 1, {}).get(category['id'], '')
                  for category in active_categories]
        if any((value or '').strip() for value in values):
            consequence_rows.append([
                str(code), _value(consequence_labels[index], 'benämning'),
                *[_value(value, f"definition {category['name']} {code}")
                  for category, value in zip(active_categories, values)],
            ])
    if consequence_rows:
        _caption(document, '5-4', 'Planerade konsekvensdefinitioner')
        headers = ['Nivå', 'Benämning'] + [category['name'] for category in active_categories]
        widths = ([16, 30] +
                  [max(35, int(204 / max(1, len(active_categories))))]
                  * len(active_categories))
        _table(document, headers, consequence_rows, widths)
    else:
        document.add_paragraph(missing('konsekvenskategorier och definitioner'))

    portrait_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _page_setup(portrait_section, landscape=False, paper='A4')
    document.add_heading('5.5 Barriärer, enablers, riskgraph och LOPA', 2)
    document.add_paragraph(
        'Om HAZOP-gruppen identifierar behov av en säkerhetsinstrumenterad funktion '
        'kan erforderlig riskreduktion bedömas med riskgraph eller LOPA. Val av '
        'metod och fördjupningsnivå görs i analysen utifrån det aktuella scenariot.')
    document.add_paragraph(
        'Befintliga och planerade barriärer ska diskuteras för det specifika '
        'scenariot. Gruppen bekräftar deras funktion, relevans och oberoende '
        'innan de kan tillgodoräknas i riskbedömningen. Enablers dokumenteras '
        'när de påverkar händelseförloppets utveckling. Dessa principer är '
        'planerade arbetssätt och innebär inte att någon barriär redan har '
        'bedömts som tillräcklig.')
    if data['field']('Barriärunderlag'):
        document.add_paragraph(data['field']('Barriärunderlag'))


def _format_frequency_per_year(value):
    """Format a catalogue frequency consistently for Swedish ToR output."""
    return f"{float(value):.6g}".replace('.', ',') + '/år'


def _add_standard_cause_catalogue_annex(document, data):
    """Add the current, frequency-backed catalogue as a one-page table annex."""
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Pt

    # The explanatory text is intentionally separate from the table. The
    # latter is a compact, landscape page that can be used directly as a
    # workshop catalogue without being split across two pages.
    introduction_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _page_setup(introduction_section, landscape=False, paper='A4')
    document.add_heading('Bilaga 2 Standardorsaker och frekvensunderlag', 1)
    document.add_paragraph(
        'Tabell B-1 är ett preliminärt katalogunderlag för HAZOP-workshopen. '
        'Frekvenserna är förslag till inledande bedömning och ska inför användning '
        'i ett enskilt scenario granskas, dokumenteras och vid behov justeras av '
        'analysgruppen med hänsyn till den aktuella anläggningen, driftmiljön och '
        'tillgängligt tillförlitlighetsunderlag. Tabellen utgör inte ett HAZOP-resultat '
        'och innebär inte att en komponent, en skyddsfunktion eller en SIF uppfyller '
        'någon angiven säkerhetsintegritetsnivå.')
    document.add_paragraph(
        'Frekvensen 0,09/år är en avrundad screeningnivå. Den har härletts från '
        'den övre gränsen för SIL 1 vid hög efterfrågan eller kontinuerligt läge, '
        'PFH < 1 × 10⁻⁵ h⁻¹ (mindre än en farlig felfunktion per 100 000 timmar). '
        'Omräkningen 1 × 10⁻⁵ h⁻¹ × 8 760 h/år ger 8,76 × 10⁻²/år och har avrundats '
        'till 0,09/år. PFH avser den genomsnittliga frekvensen av farligt fel för '
        'en specificerad säkerhetsfunktion; omräkningen får därför inte användas som '
        'en direkt komponentdatauppgift, som ett PFDavg-värde eller som verifiering '
        'av att en säkerhetsfunktion uppnår SIL 1.')
    document.add_paragraph(
        'Värdena ska särskilt omprövas när processegenskaperna avviker från '
        'antagandena i katalogen. Renare eller smutsigare medier kan exempelvis '
        'förändra sannolikheten för igensättning, korrosion, erosion, beläggning och '
        'läckage. För manuella ventiler kan manöverfrekvens, arbetssätt, åtkomlighet '
        'och mänskliga felmöjligheter vara avgörande. För elberoende utrustning kan '
        'elnätets dokumenterade tillförlitlighet, matningsarkitektur, reservkraft, '
        'spänningskvalitet och återställningsförmåga påverka den relevanta frekvensen. '
        'Projektet ska därför använda dokumenterade anläggningsdata, leverantörsdata, '
        'drifterfarenhet och tillämpliga analyser när sådant underlag finns.')

    table_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _page_setup(table_section, landscape=True, paper='A4')
    document.add_heading('Bilaga 2A Katalogtabell', 1)
    intro = document.add_paragraph(
        'Endast aktiva standardorsaker med angiven frekvens redovisas. '
        'Orsaker utan frekvens och objekttyper utan sådana orsaker har utelämnats.')
    intro.paragraph_format.space_after = Pt(2)
    _caption(document, 'B-1', 'Föreslagna standardorsaker och frekvenser')

    rows_by_object = {}
    for cause in data['standard_cause_catalogue']:
        rows_by_object.setdefault(cause['object_name'], []).append(
            f"{cause['description']} — {_format_frequency_per_year(cause['frequency'])}")
    rows = [[object_name, '\n'.join(causes)]
            for object_name, causes in rows_by_object.items()]
    table = _table(
        document,
        ['Objekt', 'Föreslagna standardorsaker och frekvenser'],
        rows or [[missing('aktiva standardorsaker med frekvens'), '']],
        [58, 215])
    # Twelve grouped rows and a small type size keep the complete, current
    # catalogue on this one landscape A4 page. The next section starts on a
    # new page, so a future longer catalogue cannot silently continue into the
    # reference text below.
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 0.9
                for run in paragraph.runs:
                    run.font.size = Pt(8 if row_index == 0 else 7.5)

    reference_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    _page_setup(reference_section, landscape=False, paper='A4')
    document.add_heading('Bilaga 2B Referensram och tillämpning', 1)
    document.add_paragraph(
        'Följande standarder utgör referensram för metod, funktionell säkerhet '
        'och den fortsatta projektanpassningen. Tillämplig utgåva, nationell '
        'adoption och eventuella kontraktskrav ska fastställas för projektet.')
    for reference in (
            'IEC 61882:2016, Hazard and operability studies (HAZOP studies) — '
            'Application guide.',
            'IEC 61511-1:2016, Functional safety — Safety instrumented systems '
            'for the process industry sector — Part 1: Framework, definitions, '
            'system, hardware and application programming requirements.',
            'IEC 61511-2:2016, Functional safety — Safety instrumented systems '
            'for the process industry sector — Part 2: Guidelines in the application '
            'of IEC 61511-1.',
            'IEC 61511-3:2017, Functional safety — Safety instrumented systems '
            'for the process industry sector — Part 3: Guidance for the determination '
            'of the required safety integrity levels.',
            'IEC 61508-1:2010, Functional safety of electrical/electronic/programmable '
            'electronic safety-related systems — Part 1: General requirements.',
            'IEC 61508-2:2010, Functional safety of electrical/electronic/programmable '
            'electronic safety-related systems — Part 2: Requirements for '
            'electrical/electronic/programmable electronic safety-related systems.'):
        document.add_paragraph(reference, style='List Bullet')
    document.add_paragraph(
        'Standarderna anger ramverk, livscykelkrav och mått för funktionell '
        'säkerhet. De ersätter inte anläggningsspecifika feldata eller en dokumenterad '
        'bedömning av den enskilda processen. Katalogvärden ska därför alltid '
        'spåras till sitt beslutade underlag när de används i riskbedömning, LOPA '
        'eller dimensionering av en säkerhetsinstrumenterad funktion.')


def build_tor(db):
    """Build a ToR from planning inputs only. Caller owns snapshot lifetime."""
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    data = collect_tor_data(db)
    field = data['field']
    project = _value(db.get_config('project_name', ''), 'system')
    client = _value(db.get_config('project_client', ''), 'kund')
    project_number = (db.get_config('project_number', '') or '').strip()
    suggested_number = f'{project_number}-ToR-01' if project_number else ''
    number = _value(field('ToRnummer') or db.get_config('tor_number', '') or suggested_number,
                    'ToR-nummer')
    latest_revision = data['revisions'][-1] if data['revisions'] else {}
    revision = _value(field('ToRrevision') or latest_revision.get('label'), 'ToR-revision')
    date = _value(field('ToRdatum') or latest_revision.get('date'), 'ToR-datum')
    prepared_by = db.get_config('tor_prepared_by', '')
    reviewed_by = db.get_config('tor_reviewed_by', '')
    approved_by = db.get_config('tor_approved_by', '')
    contact = field('Kontaktperson kund') or db.get_config('tor_contact_by', '')
    company = db.get_config('company_name', 'ProSa Process Safety Consulting AB') or 'ProSa Process Safety Consulting AB'
    company_address = field('Kontorsadress ProSa') or db.get_config('company_street', '')
    values = {
        'TITLE': 'Terms of Reference HAZOP för ' + project,
        'CLIENT': client, 'REPORT_NUMBER': number, 'REVISION': revision,
        'DATE': date, 'STATUS': revision,
        'DISTRIBUTION': _value(field('Distribution') or 'Enligt kundens anvisning', 'distribution'),
        'AUTHOR': _value(prepared_by, 'framtagen av'),
        'REVIEWER': _value(reviewed_by, 'kvalitetsgranskad av'),
        'PROSA_ADDRESS': _value(company_address, 'ProSa-adress'),
        'CLIENT_ADDRESS': field('Kundadress') or '',
        'MANAGER': field('Uppdragsansvarig') or prepared_by,
        'PROSA_CONTACT': '', 'CLIENT_PERSON': _value(contact, 'kontaktperson'),
        'CLIENT_CONTACT': field('Kontaktuppgifter kund') or '',
    }
    revisions = [{
        'REVISION': _value(row.get('label'), 'revision'),
        'REVISION_DATE': _value(row.get('date'), 'revisionsdatum'),
        'REVISION_DESCRIPTION': _value(row.get('description'), 'revisionsbeskrivning'),
        'REVISION_AUTHOR': _value(row.get('performed_by'), 'utfört av'),
    } for row in data['revisions']] or [{
        'REVISION': revision, 'REVISION_DATE': date,
        'REVISION_DESCRIPTION': 'Första utgåva',
        'REVISION_AUTHOR': _value(prepared_by, 'framtagen av'),
    }]
    document = new_report_document(values, revisions)
    document.core_properties.title = values['TITLE']
    document.core_properties.author = prepared_by or ''
    document.core_properties.subject = 'Terms of Reference för planerad HAZOP-analys'
    _bookmark_cover_value(document, date, _TOR_DATE_BOOKMARK)
    _bookmark_cover_value(document, number, _TOR_NUMBER_BOOKMARK)
    source_header = document.sections[-1].header.part
    source_elements = deepcopy(list(source_header._element))
    source_relationships = tuple(source_header.rels.items())
    for section in document.sections:
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.text = (run.text.replace('HAZOP-rapport', 'HAZOP ToR')
                                        .replace('Report', 'ToR'))

    _front_heading(document, 'Syfte med dokumentet')
    document.add_paragraph(
        'Detta Terms of Reference beskriver förutsättningarna för den planerade '
        f'HAZOP-analysen av {project}. Dokumentet är ett arbetsunderlag före '
        'analysen och redovisar därför inte HAZOP-resultat, riskutfall eller rekommendationer.')
    document.add_page_break()
    document.add_paragraph('Innehåll', 'TOC Heading')
    _field_run(document.add_paragraph(), 'TOC \\o "1-2" \\h \\z',
               'Uppdatera innehållsförteckningen i Word med Ctrl+A och F9.')
    document.add_section(WD_SECTION_START.CONTINUOUS)

    document.add_heading('1 Inledning', 1)
    document.add_paragraph(
        f'{client} planerar en HAZOP-analys av {project}. {company} ska leda '
        'och dokumentera workshopen. HAZOP används för att systematiskt utmana '
        'designavsikter, gränssnitt och relevanta driftfall före eller under projektets fortsatta arbete.')
    document.add_heading('1.1 Syfte och omfattning', 2)
    document.add_paragraph(field('Syfte') or
        'Syftet är att bekräfta och utmana systemets utformning, identifiera '
        'risker och operabilitetsfrågor samt besluta om eventuellt behov av fortsatt utredning.')
    document.add_paragraph(field('Omfattning') or
        'Omfattningen utgår från den preliminära nodindelningen och de dokument '
        'som listas i detta ToR.')
    document.add_heading('1.2 Avgränsningar och förutsättningar', 2)
    document.add_paragraph(field('Avgränsningar') or
        'Analysen behandlar de systemdelar och gränssnitt som bekräftas av HAZOP-gruppen. '
        'Nodindelning, dokumentunderlag och planerade driftfall kan justeras under workshopen.')
    if field('Driftfall'):
        document.add_paragraph('Planerade driftfall: ' + field('Driftfall'))
    if field('Analysförutsättningar'):
        document.add_paragraph(field('Analysförutsättningar'))

    _chapter(document, '2 Dokumentunderlag')
    document.add_paragraph(
        'P&ID-ritningar och övrigt dokumentunderlag används för att bekräfta '
        'nodgränser, designavsikter och relevanta gränssnitt före analysen.')
    _caption(document, '2-1', 'Registrerade P&ID-ritningar')
    _table(document, ['Dokumentnummer', 'Dokumenttitel', 'Revision', 'Datum', 'PDF-sida'], [
        [_value(sheet.get('drawing_number'), 'ritningsnummer'),
         _value(sheet.get('drawing_name') or sheet.get('sheet_name'), 'ritningsnamn'),
         _value(sheet.get('drawing_revision'), 'revision'),
         _value(sheet.get('drawing_date'), 'ritningsdatum'), str(sheet['physical_page'] + 1)]
        for sheet in data['sheets']] or [[missing('P&ID-ritning'), '', '', '', '']],
        [34, 51, 23, 32, 20])
    if field('Övriga referensdokument'):
        document.add_heading('2.1 Övriga referensdokument', 2)
        for text in field('Övriga referensdokument').split('\n'):
            document.add_paragraph(text)

    _chapter(document, '3 HAZOP-deltagare och planering')
    document.add_paragraph(
        'HAZOP-workshopen genomförs med en tvärdisciplinär expertgrupp. '
        'Deltagaruppgifterna är preliminära tills analysen är genomförd.')
    _add_participant_table(document, data)
    _add_planned_sessions(document, data)

    _chapter(document, '4 HAZOP-metodik och planeringsunderlag')
    document.add_paragraph(
        'HAZOP-metoden använder parametrar och guideord för att identifiera '
        'meningsfulla avvikelser från respektive nods designavsikt. Varje relevant '
        'avvikelse behandlas av gruppen med avseende på möjliga orsaker, konsekvenser '
        'och befintliga eller planerade skydd.')
    document.add_heading('4.1 Guideord och avvikelser', 2)
    document.add_paragraph(
        'Tabell 4-1 visar de guideord och avvikelser som är aktuella i projektets '
        'bibliotek. Guideorden kan justeras, kompletteras eller begränsas när '
        'nodernas designavsikter och driftfall bekräftas under workshopen.')
    _caption(document, '4-1', 'Guideord och avvikelser')
    _table(document, ['Guideord eller avvikelse'], [
        [_value(word.get('description'), 'guideord eller avvikelse')]
        for word in data['guide_words']] or [[missing('guideord eller avvikelser')]], [160])
    _add_nodes(document, db, data)

    _add_risk_framework(document, db, data)
    document.add_heading('Bilaga 1 Förberedelsechecklista', 1)
    document.add_paragraph(
        'Före analysen ska aktuella P&ID-ritningar, systembeskrivningar, '
        'designavsikter, driftfall, nodgränser och deltagarlista bekräftas. '
        'Förändringar dokumenteras i projektets förberedelseunderlag.')
    _add_standard_cause_catalogue_annex(document, data)

    _copy_running_headers(document, source_elements, source_relationships)
    footer = document.sections[2].footer
    for child in list(footer._element):
        footer._element.remove(child)
    footer._element.append(OxmlElement('w:p'))
    paragraph = footer.paragraphs[0]
    paragraph.alignment = 2
    paragraph.add_run('Sida ')
    _field_run(paragraph, 'PAGE')
    paragraph.add_run(' av ')
    _field_run(paragraph, 'NUMPAGES')
    page_number_type = OxmlElement('w:pgNumType')
    page_number_type.set(qn('w:start'), '1')
    document.sections[2]._sectPr.append(page_number_type)
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    document.settings.element.append(update)
    apply_report_fonts(document)
    _highlight_document(document)
    return document


def export_tor_word(db, filepath):
    """Export atomically using the same read-only contract as report export."""
    target = Path(filepath)
    temporary = None
    try:
        with _report_snapshot(db) as snapshot:
            document = build_tor(snapshot)
        with tempfile.NamedTemporaryFile(dir=target.parent, suffix='.docx', delete=False) as file:
            temporary = Path(file.name)
        document.save(temporary)
        os.replace(temporary, target)
        return True, ''
    except ImportError:
        return False, 'python-docx saknas.\\nKör: pip install python-docx'
    except Exception as error:
        return False, str(error)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
